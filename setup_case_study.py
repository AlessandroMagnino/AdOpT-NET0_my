import h5py
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pathlib import Path
import os
import json

def setup_case_study(pathway: str, year: str):
    """
    Check if carriers files are ok
    Define technologies.xlsx
    """

    case_study_path = Path("case_studies") / pathway / year

    ### Carriers check
    carriers = pd.read_excel(case_study_path / "carriers_list.xlsx", sheet_name="carriers")

    for carrier in carriers["CARRIERS"]:
        # Check if carrier file exists
        carrier_file = case_study_path / "carriers" / f"{carrier}.xlsx"
        if not os.path.exists(carrier_file):
            raise FileNotFoundError(f"Carrier file {carrier_file} not found for carrier {carrier} in year {year}. Please check the file path and ensure it exists.")
        
    ### Technologies definition
    couples = {
        "2025": "2020",
        "2030": "2025",
        "2040": "2030",
        "2050": "2040",
    }

    em_limits = {
        "2040": 0.5,
        "2050": 0.0
    }

    if year not in couples:
        raise ValueError(f"No previous year configured for {year}. Update the couples map in setup_case_study.py.")

    previous_year = couples[year]

    xlxs_ref = Path("data") / "technologies_fac_simile.xlsx" # just to have a reference file to copy
    output_path  = case_study_path / "technologies.xlsx"
    data = Path("output") / pathway / previous_year / "optimization_results.h5"

    if not data.exists():
        raise FileNotFoundError(f"Previous year results not found: {data}")

    base = "design/nodes/2022"
    rows = []
    with h5py.File(data, "r") as f:
        for node in f[base].keys():
            for comp in f[f"{base}/{node}"].keys():
                ds_path = f"{base}/{node}/{comp}/size"
                size = _to_scalar(f[ds_path][()]) if ds_path in f else 0.0
                rows.append({"node": node, "component": comp, "size": size})

    pivot = pd.DataFrame(rows).pivot(index="node", columns="component", values="size").fillna(0)

    # Read reference technologies and create output table
    wb_ref    = load_workbook(xlxs_ref, read_only=True)
    ws_ref    = wb_ref["existing"]
    header_ref = list(next(ws_ref.iter_rows(min_row=1, max_row=1, values_only=True)))
    nodes_ref  = [row[0] for row in ws_ref.iter_rows(min_row=2, values_only=True)]
    tech_cols  = header_ref[1:]

    # Create output DataFrame with all nodes and technologies, filling missing values with 0
    output_rows = []
    for node in nodes_ref:
        row = {"cluster": node}
        for tech in tech_cols:
            if tech in pivot.columns and node in pivot.index:
                row[tech] = round(pivot.loc[node, tech], 2)
            else:
                row[tech] = 0
        output_rows.append(row)

    out_df = pd.DataFrame(output_rows)

    # Write formatted excel output (existing + new)
    wb = Workbook()

    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill("solid", start_color="2E4057")
    data_font = Font(name="Arial", size=10)
    alt_fill = PatternFill("solid", start_color="EAF0FB")
    white_fill = PatternFill("solid", start_color="FFFFFF")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    all_cols = ["cluster"] + tech_cols

    def write_sheet(ws, df):
        for c_idx, col_name in enumerate(all_cols, start=1):
            cell = ws.cell(row=1, column=c_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = border

        for r_idx, row in enumerate(df.itertuples(index=False), start=2):
            alt = r_idx % 2 == 0
            for c_idx, val in enumerate(row, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                cell.font = data_font
                cell.fill = alt_fill if alt else white_fill
                cell.alignment = left_align if c_idx == 1 else center_align
                cell.border = border
                if c_idx > 1:
                    cell.number_format = "#,##0.00;-#,##0.00;0.00"

        ws.column_dimensions["A"].width = 12
        for c_idx in range(2, len(all_cols) + 1):
            ws.column_dimensions[get_column_letter(c_idx)].width = max(
                len(all_cols[c_idx - 1]) + 2, 10
            )

        ws.row_dimensions[1].height = 18
        ws.freeze_panes = "B2"

    ws_existing = wb.active
    ws_existing.title = "existing"
    write_sheet(ws_existing, out_df)

    out_df_new = out_df.copy()
    out_df_new.loc[:, tech_cols] = 1

    ws_new = wb.create_sheet(title="new")
    write_sheet(ws_new, out_df_new)

    ws_existing, ws_new = electricity_connection_correction(ws_existing, ws_new)

    ws_existing = H2_network_connection_correction(ws_existing)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    # Set limits on emissions
    if year in em_limits:
        # Check emissions in previous year results
        with h5py.File(data, "r") as f:
            # Go in /summary/emissions_net and take value
            em_prev_year = _to_scalar(f["summary/emissions_net"][()])
            em_limit = em_prev_year * em_limits[year]
    
        # Save value in config file
        config_path = case_study_path / "config_specs.json"
        if not config_path.exists():
            raise FileNotFoundError(f"Config file not found: {config_path}")
        
        config_specs = json.loads(config_path.read_text())
        config_specs["optimization"]["objective"]["value"] = "costs_emissionlimit"
        config_specs["optimization"]["emission_limit"]["value"] = em_limit

        with open(config_path, "w") as f:
            json.dump(config_specs, f)

    return


##### Read sizes from h5 and create formatted excel output
def _to_scalar(x):
    x = np.asarray(x)
    return float(x.reshape(-1)[0]) if x.size == 1 else float("nan")


def electricity_connection_correction(ws_existing, ws_new):
    """
    Just to manually correct data on electricity connections
    """
    
    electricity_connections = [
        "electricity_grid_connection_BE",
        "electricity_grid_connection_DE",
        "electricity_grid_connection_NL"
    ]

    couples_node_electricity_connections = {
        "BEL1": ["electricity_grid_connection_BE"],
        "BEL2": ["electricity_grid_connection_BE"],
        "DE": ["electricity_grid_connection_DE"],
        "NL1": ["electricity_grid_connection_NL"],
        "NL2": ["electricity_grid_connection_NL"],
        "NL3": ["electricity_grid_connection_NL"],
        "NL4": ["electricity_grid_connection_NL"]
        }

    # All values in existing sheet and columns of electricity connections to 0
    for row in ws_existing.iter_rows(min_row=2, min_col=2, max_col=ws_existing.max_column):
        for cell in row:
            if cell.column_letter in [get_column_letter(i) for i, tech in enumerate(ws_existing[1], start=1) if tech.value in electricity_connections]:
                cell.value = 0

    # In new sheet, couple nodes and electricity connections, set value to 1
    for row in ws_new.iter_rows(min_row=2, min_col=1, max_col=ws_new.max_column):
        node_cell = row[0]
        for cell in row[1:]:
            tech = ws_new.cell(row=1, column=cell.column).value
            if tech in electricity_connections and tech in couples_node_electricity_connections.get(node_cell.value, []):
                cell.value = 1
            elif tech in electricity_connections:
                cell.value = 0

    return ws_existing, ws_new


def H2_network_connection_correction(ws_existing):
    """
    Just to manually correct data on H2 connections
    """
    
    H2_connections = [
        "H2_network_connection_in",
        "H2_network_connection_out"
    ]

    # In existing sheet, set all values in columns of H2 connections to 0
    for row in ws_existing.iter_rows(min_row=2, min_col=2, max_col=ws_existing.max_column):
        for cell in row:
            if cell.column_letter in [get_column_letter(i) for i, tech in enumerate(ws_existing[1], start=1) if tech.value in H2_connections]:
                cell.value = 0

    return ws_existing